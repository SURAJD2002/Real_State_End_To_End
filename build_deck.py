#!/usr/bin/env python3
"""
Planwise Enterprise — Business / Partnership Pitch Deck Generator
Creates a 16-slide widescreen (16:9) executive presentation.
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# ---------------------------------------------------------
# CONSTANTS & THEME
# ---------------------------------------------------------
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Color Palette (Deep Charcoal + Warm White + Planwise Blue)
COLOR_BG_DARK = RGBColor(15, 20, 28)       # #0F141C
COLOR_CARD_BG = RGBColor(24, 32, 47)       # #18202F
COLOR_CARD_ALT = RGBColor(19, 26, 38)      # #131A26
COLOR_CARD_HIGHLIGHT = RGBColor(30, 41, 59)# #1E293B
COLOR_BORDER = RGBColor(45, 55, 72)        # #2D3748
COLOR_BORDER_ACCENT = RGBColor(56, 189, 248) # #38BDF8
COLOR_BORDER_BLUE = RGBColor(37, 99, 235)  # #2563EB

COLOR_TEXT_TITLE = RGBColor(248, 250, 252) # #F8FAFC
COLOR_TEXT_BODY = RGBColor(226, 232, 240)  # #E2E8F0
COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
COLOR_TEXT_SUBTLE = RGBColor(100, 116, 139)# #64748B

COLOR_ACCENT_BLUE = RGBColor(37, 99, 235)  # #2563EB
COLOR_ACCENT_CYAN = RGBColor(56, 189, 248) # #38BDF8
COLOR_ACCENT_EMERALD = RGBColor(16, 185, 129) # #10B981
COLOR_ACCENT_AMBER = RGBColor(245, 158, 11)   # #F59E0B

FONT_HEADING = "Arial"
FONT_BODY = "Arial"

ASSETS_DIR = os.path.abspath("assets/pitch_deck")


def create_base_slide(prs, category, title, subtitle, slide_num, total_slides=16):
    """Creates a blank slide with dark theme, category chip, title, subtitle, and footer."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 1. Background fill
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG_DARK
    bg.line.fill.background()

    # 2. Category badge / chip
    if category:
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.48), Inches(8.0), Inches(0.35))
        tf = cat_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.name = FONT_HEADING
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_CYAN

    # 3. Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.82), Inches(11.733), Inches(0.55))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(23)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_TITLE

    # 4. Subtitle / Supporting Sentence
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.4))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(12.5)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED

    # 5. Footer divider & branding
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.92), Inches(11.733), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_BORDER
    line.line.fill.background()

    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.98), Inches(9.0), Inches(0.3))
    tf_f = footer_box.text_frame
    tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
    pf = tf_f.paragraphs[0]
    pf.text = "PLANWISE ENTERPRISE  •  Land-to-Home Development & Construction Platform"
    pf.font.name = FONT_BODY
    pf.font.size = Pt(8.5)
    pf.font.color.rgb = COLOR_TEXT_SUBTLE

    num_box = slide.shapes.add_textbox(Inches(10.533), Inches(6.98), Inches(2.0), Inches(0.3))
    tf_n = num_box.text_frame
    tf_n.margin_left = tf_n.margin_top = tf_n.margin_right = tf_n.margin_bottom = 0
    pn = tf_n.paragraphs[0]
    pn.alignment = PP_ALIGN.RIGHT
    pn.text = f"{slide_num} / {total_slides}"
    pn.font.name = FONT_BODY
    pn.font.size = Pt(8.5)
    pn.font.color.rgb = COLOR_TEXT_SUBTLE

    return slide


def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
    """Creates a rounded card container."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card


# =========================================================
# SLIDE BUILDERS
# =========================================================

def build_slide_1(prs):
    """Slide 1: Cover"""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG_DARK
    bg.line.fill.background()

    # Left content area (x=0.8, y=1.2, w=5.8, h=5.2)
    # Category Pill
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.1), Inches(2.8), Inches(0.38))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(23, 37, 84) # Dark blue
    pill.line.color.rgb = COLOR_ACCENT_BLUE
    pill.line.width = Pt(1)
    tf_pill = pill.text_frame
    tf_pill.margin_left = tf_pill.margin_top = tf_pill.margin_right = tf_pill.margin_bottom = 0
    p = tf_pill.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "PLANWISE ENTERPRISE"
    p.font.name = FONT_HEADING
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_CYAN

    # Sub-badge
    badge_box = slide.shapes.add_textbox(Inches(3.75), Inches(1.12), Inches(3.0), Inches(0.35))
    tf_b = badge_box.text_frame
    tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
    pb = tf_b.paragraphs[0]
    pb.text = "Computational Development Platform"
    pb.font.name = FONT_BODY
    pb.font.size = Pt(10)
    pb.font.color.rgb = COLOR_TEXT_MUTED

    # Main Headline
    head_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.75), Inches(5.8), Inches(2.2))
    tf_h = head_box.text_frame
    tf_h.word_wrap = True
    tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0
    ph = tf_h.paragraphs[0]
    ph.text = "From a piece of land to a build-ready project."
    ph.font.name = FONT_HEADING
    ph.font.size = Pt(36)
    ph.font.bold = True
    ph.font.color.rgb = COLOR_TEXT_TITLE

    # Platform Descriptor
    desc_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.05), Inches(5.8), Inches(0.5))
    tf_d = desc_box.text_frame
    tf_d.margin_left = tf_d.margin_top = tf_d.margin_right = tf_d.margin_bottom = 0
    pd = tf_d.paragraphs[0]
    pd.text = "Land-to-Home Development & Construction Platform"
    pd.font.name = FONT_HEADING
    pd.font.size = Pt(15)
    pd.font.bold = True
    pd.font.color.rgb = COLOR_ACCENT_CYAN

    # Supporting paragraph
    body_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.65), Inches(5.6), Inches(1.2))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    tf_body.margin_left = tf_body.margin_top = tf_body.margin_right = tf_body.margin_bottom = 0
    pbody = tf_body.paragraphs[0]
    pbody.text = "Unifying statutory land feasibility, computational house design, model-linked quantities, regional cost waterfalls, and licensed professional review into one traceable, deterministic workflow."
    pbody.font.name = FONT_BODY
    pbody.font.size = Pt(12)
    pbody.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom notice
    foot_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(5.6), Inches(0.4))
    tf_foot = foot_box.text_frame
    pfoot = tf_foot.paragraphs[0]
    pfoot.text = "Confidential Business Presentation  •  October 2026"
    pfoot.font.name = FONT_BODY
    pfoot.font.size = Pt(9.5)
    pfoot.font.color.rgb = COLOR_TEXT_SUBTLE

    # Right Side Architectural Visual
    cover_img_path = os.path.join(ASSETS_DIR, "cover_hero.jpg")
    if os.path.exists(cover_img_path):
        slide.shapes.add_picture(cover_img_path, Inches(6.85), Inches(1.0), Inches(5.683), Inches(5.5))
        # Add a clean crisp rectangular border frame
        frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.85), Inches(1.0), Inches(5.683), Inches(5.5))
        frame.fill.background()
        frame.line.color.rgb = COLOR_BORDER
        frame.line.width = Pt(1.5)


def build_slide_2(prs):
    """Slide 2: The Problem"""
    slide = create_base_slide(prs, "Market Reality", "Real estate development is fragmented.",
                             "A single residential parcel passes through disconnected silos, causing repeated rework, cost drift, and project delays.", 2)

    # Top Sequential Pipeline Box
    top_box = add_card(slide, Inches(0.8), Inches(1.95), Inches(11.733), Inches(1.15), bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)
    tb_text = slide.shapes.add_textbox(Inches(1.0), Inches(2.05), Inches(11.333), Inches(0.95))
    tf = tb_text.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p0 = tf.paragraphs[0]
    p0.text = "THE TRADITIONAL SEQUENTIAL RELAY (HIGH FRICTION & DATA RE-CREATION)"
    p0.font.name = FONT_HEADING
    p0.font.size = Pt(9.5)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT_AMBER

    p1 = tf.add_paragraph()
    p1.text = "Land  →  Survey  →  Feasibility  →  Architect  →  Design  →  BOQ  →  Costing  →  Engineer  →  Approvals  →  Contractor  →  Construction"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_TITLE

    p2 = tf.add_paragraph()
    p2.text = "⚠ Every handoff occurs between disjointed software tools, PDFs, and spreadsheets with zero shared single source of truth."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COLOR_TEXT_MUTED

    # 3 Column Structural Problem Cards
    col_w = Inches(3.68)
    gap = Inches(0.34)
    y_card = Inches(3.3)
    h_card = Inches(3.35)

    problems = [
        ("01", "Disconnected Data Silos",
         "Data gets recreated repeatedly",
         [
             "Surveys, CAD drawings, and spreadsheets live in isolated files across separate consulting firms.",
             "Every stakeholder manually redraws project geometry from scratch, introducing discrepancies.",
             "No single traceable project model connects land coordinates to construction handoff."
         ]),
        ("02", "Rework & Cost Drift",
         "Design adjustments break downstream cost",
         [
             "Statutory feasibility is largely manual and error-prone during early land evaluation.",
             "Downstream BOQ and cost estimates drift continuously from iterative design modifications.",
             "Cost blindspots and scope mismatches are discovered too late, triggering painful redesigns."
         ]),
        ("03", "Delayed Validation",
         "Professional review happens too late",
         [
             "Licensed structural and statutory compliance reviews only occur after design freeze.",
             "Non-compliant designs require costly re-drafting and delay municipal approval cycles.",
             "Customers and builders struggle to interpret opaque technical and engineering decisions."
         ])
    ]

    for i, (num, title, subtitle, bullets) in enumerate(problems):
        x = Inches(0.8) + i * (col_w + gap)
        add_card(slide, x, y_card, col_w, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.25), y_card + Inches(0.2), col_w - Inches(0.5), h_card - Inches(0.4))
        tf_c = t_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0

        p_num = tf_c.paragraphs[0]
        p_num.text = f"PROBLEM {num}"
        p_num.font.name = FONT_HEADING
        p_num.font.size = Pt(9)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_ACCENT_CYAN

        p_t = tf_c.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_TITLE

        p_sub = tf_c.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(10)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_ACCENT_AMBER

        tf_c.add_paragraph() # space

        for bullet in bullets:
            pb = tf_c.add_paragraph()
            pb.text = f"•  {bullet}"
            pb.font.name = FONT_BODY
            pb.font.size = Pt(10)
            pb.font.color.rgb = COLOR_TEXT_MUTED


def build_slide_3(prs):
    """Slide 3: The Planwise Solution"""
    slide = create_base_slide(prs, "Core Innovation", "One computational workflow. One source of truth.",
                             "The Canonical Building Model (CBM) links every stage from land parcel coordinates to licensed build release.", 3)

    # Left Side: 8-Stage Unified Workflow (x=0.8, w=5.5, y=1.95, h=4.7)
    left_card = add_card(slide, Inches(0.8), Inches(1.95), Inches(5.5), Inches(4.75), bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

    t_box = slide.shapes.add_textbox(Inches(1.05), Inches(2.1), Inches(5.0), Inches(4.45))
    tf = t_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_head = tf.paragraphs[0]
    p_head.text = "THE UNIFIED PLANWISE WORKFLOW"
    p_head.font.name = FONT_HEADING
    p_head.font.size = Pt(10)
    p_head.font.bold = True
    p_head.font.color.rgb = COLOR_ACCENT_CYAN

    stages = [
        ("01", "Land Intake", "SECS topocentric frame & boundary survey"),
        ("02", "Statutory Feasibility", "Mumbai DCPR 2034 rules, setbacks & envelope"),
        ("03", "House Generation", "Multi-constraint CP-SAT parametric layouts"),
        ("04", "Canonical Building Model", "Semantic 3D single source of truth (CBM)"),
        ("05", "Model-Linked QTO", "Automated IS 1200 rules & geometric deductions"),
        ("06", "Cost Waterfall", "Regional trade, material, labour & margin breakdown"),
        ("07", "Professional Review", "G0–G6 multi-discipline engineer review ledger"),
        ("08", "Build Release", "Cryptographic SHA-256 seal & IFC4 BIM package")
    ]

    for num, name, detail in stages:
        p_item = tf.add_paragraph()
        p_item.text = f"{num}  {name}"
        p_item.font.name = FONT_HEADING
        p_item.font.size = Pt(11)
        p_item.font.bold = True
        if "Canonical" in name:
            p_item.font.color.rgb = COLOR_ACCENT_CYAN
        else:
            p_item.font.color.rgb = COLOR_TEXT_TITLE

        p_det = tf.add_paragraph()
        p_det.text = f"     {detail}"
        p_det.font.name = FONT_BODY
        p_det.font.size = Pt(9.5)
        p_det.font.color.rgb = COLOR_TEXT_MUTED

    # Right Side: The CBM Differentiator (x=6.55, w=5.98, y=1.95, h=4.75)
    right_card = add_card(slide, Inches(6.55), Inches(1.95), Inches(5.98), Inches(4.75), bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER_BLUE)

    cbm_img_path = os.path.join(ASSETS_DIR, "cbm_bim.jpg")
    if os.path.exists(cbm_img_path):
        slide.shapes.add_picture(cbm_img_path, Inches(6.75), Inches(2.1), Inches(5.58), Inches(2.2))

    t_box_r = slide.shapes.add_textbox(Inches(6.75), Inches(4.45), Inches(5.58), Inches(2.1))
    tf_r = t_box_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0

    pr_head = tf_r.paragraphs[0]
    pr_head.text = "THE CANONICAL BUILDING MODEL (CBM) DIFFERENTIATOR"
    pr_head.font.name = FONT_HEADING
    pr_head.font.size = Pt(10.5)
    pr_head.font.bold = True
    pr_head.font.color.rgb = COLOR_ACCENT_CYAN

    pr_desc = tf_r.add_paragraph()
    pr_desc.text = "The same underlying mathematical building model simultaneously drives:"
    pr_desc.font.name = FONT_BODY
    pr_desc.font.size = Pt(10)
    pr_desc.font.color.rgb = COLOR_TEXT_MUTED

    points = [
        "Floor plans & 2D/3D geometry without manual drafting",
        "Automated 12-point statutory & physical geometric validation",
        "Model-linked material takeoff (concrete, AAC blocks, plaster)",
        "Hierarchical 7-trade BOQ tied directly to 3D element IDs",
        "Live regional construction cost waterfalls",
        "Multi-discipline licensed engineer review and tamper-evident release"
    ]

    for pt in points:
        p_pt = tf_r.add_paragraph()
        p_pt.text = f"✓  {pt}"
        p_pt.font.name = FONT_BODY
        p_pt.font.size = Pt(9.5)
        p_pt.font.color.rgb = COLOR_TEXT_BODY


def build_slide_4(prs):
    """Slide 4: Product Modules"""
    slide = create_base_slide(prs, "Platform Modules", "What Planwise actually does",
                             "Six integrated computational modules replacing weeks of fragmented pre-construction consulting.", 4)

    modules = [
        ("01", "Land Intelligence",
         "Reconstructs surveyed parcel boundaries into an authoritative topocentric SECS coordinate frame with geodetic GPS anchoring."),
        ("02", "Feasibility Engine",
         "Applies statutory municipal rulepacks (e.g. Mumbai DCPR 2034) to calculate road widening, setbacks, FSI, and buildable envelopes in seconds."),
        ("03", "Computational Design",
         "Solves multi-constraint topological layouts using Google OR-Tools CP-SAT to produce Pareto-optimal, compliant residential floor plans."),
        ("04", "Quantity & Cost Engine",
         "Extracts exact material volumes directly from 3D geometry using IS 1200 deduction rules, linked to regional trade rate snapshots."),
        ("05", "Professional Review",
         "Enforces a structured G0–G6 verification ledger for certified land surveyors, architects, and structural engineers."),
        ("06", "Build Handoff",
         "Compiles tamper-evident IFC4 BIM files, validation audit reports, itemized BOQs, and cryptographic SHA-256 release certificates.")
    ]

    w_card = Inches(3.68)
    gap = Inches(0.34)
    h_card = Inches(2.2)

    for i, (num, title, desc) in enumerate(modules):
        col = i % 3
        row = i // 3
        x = Inches(0.8) + col * (w_card + gap)
        y = Inches(2.05) + row * (h_card + Inches(0.3))

        add_card(slide, x, y, w_card, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.25), y + Inches(0.2), w_card - Inches(0.5), h_card - Inches(0.4))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_num = tf.paragraphs[0]
        p_num.text = f"MODULE {num}"
        p_num.font.name = FONT_HEADING
        p_num.font.size = Pt(9.5)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_ACCENT_CYAN

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_TITLE

        tf.add_paragraph() # space

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = COLOR_TEXT_MUTED


def build_slide_5(prs):
    """Slide 5: Customer Experience"""
    slide = create_base_slide(prs, "Customer Journey", "Complex engineering underneath. Simple experience on top.",
                             "Landowners, builders, and developers navigate a clear, transparent decision journey without technical friction.", 5)

    steps = [
        ("1", "Tell us about your land",
         "Input coordinates, CTS revenue number, or upload registered surveyor boundary sketch."),
        ("2", "See what can be built",
         "Instant visual statutory envelope, permissible FSI limits, and municipal setback overlay."),
        ("3", "Choose your house",
         "Explore multiple Pareto-optimal floor plans and 3D architectural massing alternatives."),
        ("4", "See estimated cost",
         "Live, transparent preliminary cost breakdown by trades, materials, labour, and prelims."),
        ("5", "Request review",
         "Engage certified engineers and surveyors for structured multi-discipline gate verification."),
        ("6", "Move to build",
         "Receive stamped engineering audit, IFC4 BIM model, and contractor-ready 7-trade BOQ.")
    ]

    w_step = Inches(1.8)
    gap = Inches(0.186)
    y_step = Inches(2.05)
    h_step = Inches(3.6)

    for i, (num, title, desc) in enumerate(steps):
        x = Inches(0.8) + i * (w_step + gap)
        add_card(slide, x, y_step, w_step, h_step, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.15), y_step + Inches(0.15), w_step - Inches(0.3), h_step - Inches(0.3))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_step = tf.paragraphs[0]
        p_step.text = f"STEP 0{num}"
        p_step.font.name = FONT_HEADING
        p_step.font.size = Pt(9.5)
        p_step.font.bold = True
        p_step.font.color.rgb = COLOR_ACCENT_CYAN

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_TITLE

        tf.add_paragraph()

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom emphasis banner
    banner = add_card(slide, Inches(0.8), Inches(5.85), Inches(11.733), Inches(0.85), bg_color=RGBColor(23, 37, 84), border_color=COLOR_ACCENT_BLUE)
    t_banner = slide.shapes.add_textbox(Inches(1.0), Inches(5.95), Inches(11.333), Inches(0.65))
    tf_b = t_banner.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
    pb1 = tf_b.paragraphs[0]
    pb1.text = "THE USER EXPERIENCE PROMISE"
    pb1.font.name = FONT_HEADING
    pb1.font.size = Pt(9)
    pb1.font.bold = True
    pb1.font.color.rgb = COLOR_ACCENT_CYAN
    pb2 = tf_b.add_paragraph()
    pb2.text = "Replaces 4–8 weeks of fragmented consulting with minutes of computational clarity, giving property owners total transparency before committing capital."
    pb2.font.name = FONT_BODY
    pb2.font.size = Pt(10.5)
    pb2.font.color.rgb = COLOR_TEXT_TITLE


def build_slide_6(prs):
    """Slide 6: What Is Already Built"""
    slide = create_base_slide(prs, "Verified Capabilities", "This is already a working platform.",
                             "Planwise is not an unbuilt concept—the core end-to-end computational pipeline is fully implemented.", 6)

    col_w = Inches(3.68)
    gap = Inches(0.34)
    y_card = Inches(2.05)
    h_card = Inches(4.35)

    columns = [
        ("Land, Law & Geometry",
         COLOR_ACCENT_CYAN,
         [
             "SECS geodetic-to-planar coordinate transformation",
             "Site evidence engine with surveyor licence tracking",
             "Mumbai DCPR 2034 statutory rule engine (road widening & setbacks)",
             "Customer design brief intake (hard & soft constraints)",
             "CP-SAT parametric layout generator (Google OR-Tools)",
             "Canonical Building Model (CBM) semantic 3D data structure",
             "12 automated statutory & physical geometric validation checks"
         ]),
        ("Quantities, Cost & Verification",
         COLOR_ACCENT_BLUE,
         [
             "Model-linked QTO takeoff with IS 1200 deduction rules",
             "Hierarchical 7-trade BOQ compilation",
             "Regional cost waterfall engine (INDIA-MUMBAI-2026-Q4-V1)",
             "Build Review customer acknowledgement workflow",
             "Engineer Review Ledger with G0–G6 multi-discipline gates",
             "Cryptographic SHA-256 release fingerprinting",
             "Export to IFC4 BIM, JSON manifests, and CSV takeoffs"
         ]),
        ("Operational Evidence & UI",
         COLOR_ACCENT_EMERALD,
         [
             "Full-stack application: FastAPI gateway + React/TypeScript UI",
             "Interactive CAD parcel workspace with snapping & ortho tools",
             "Mapbox GIS workspace with cadastral overlays",
             "131 backend unit & regression tests passing",
             "20 Vitest geometry engine tests passing",
             "PostgreSQL schema with immutable release ledgers",
             "Controlled single-project pilot executed end-to-end"
         ])
    ]

    for i, (title, color, items) in enumerate(columns):
        x = Inches(0.8) + i * (col_w + gap)
        add_card(slide, x, y_card, col_w, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.2), y_card + Inches(0.2), col_w - Inches(0.4), h_card - Inches(0.4))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        tf.add_paragraph() # space

        for item in items:
            pi = tf.add_paragraph()
            pi.text = f"✓  {item}"
            pi.font.name = FONT_BODY
            pi.font.size = Pt(9.5)
            pi.font.color.rgb = COLOR_TEXT_BODY

    # Notice at bottom of slide
    n_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.35))
    tf_n = n_box.text_frame
    pn = tf_n.paragraphs[0]
    pn.text = "*Note: While core feasibility, generation, QTO, costing, and review are fully functional, physical construction authorization remains strictly gated for licensed professionals."
    pn.font.name = FONT_BODY
    pn.font.size = Pt(8.5)
    pn.font.color.rgb = COLOR_TEXT_SUBTLE


def build_slide_7(prs):
    """Slide 7: Real Project / Pilot"""
    slide = create_base_slide(prs, "Pilot Verification", "From software to a controlled real-project workflow.",
                             "Executed end-to-end on an authoritative 1,100 sq ft residential parcel in Bandra West, Mumbai.", 7)

    # 4 Stat Cards
    stat_w = Inches(2.78)
    stat_gap = Inches(0.2)
    y_stat = Inches(2.05)
    h_stat = Inches(1.1)

    stats = [
        ("AUTHORITATIVE PLOT AREA", "1,100.0 sq ft", "102.19 m² preserved invariant"),
        ("LOCATION & RECORD", "Bandra West, Mumbai", "CTS-1842-BANDRA (DCPR 2034)"),
        ("SELECTED DESIGN OPTION", "Option B (G+1)", "98.16 m² Gross BUA (76.8 m² Carpet)"),
        ("PRELIMINARY COST WATERFALL", "₹13,18,391 INR", "Direct preliminary construction cost")
    ]

    for i, (label, val, sub) in enumerate(stats):
        x = Inches(0.8) + i * (stat_w + stat_gap)
        add_card(slide, x, y_stat, stat_w, h_stat, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.15), y_stat + Inches(0.12), stat_w - Inches(0.3), h_stat - Inches(0.24))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = label
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(8)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_ACCENT_CYAN

        p1 = tf.add_paragraph()
        p1.text = val
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_TITLE

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_BODY
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = COLOR_TEXT_MUTED

    # 2 Big Columns Below
    col2_w = Inches(5.72)
    gap2 = Inches(0.293)
    y_col2 = Inches(3.3)
    h_col2 = Inches(3.1)

    # Column 1: 12-Stage Trace
    add_card(slide, Inches(0.8), y_col2, col2_w, h_col2, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)
    t_trace = slide.shapes.add_textbox(Inches(1.0), y_col2 + Inches(0.15), col2_w - Inches(0.4), h_col2 - Inches(0.3))
    tf_tr = t_trace.text_frame
    tf_tr.word_wrap = True
    tf_tr.margin_left = tf_tr.margin_top = tf_tr.margin_right = tf_tr.margin_bottom = 0

    p_tr_head = tf_tr.paragraphs[0]
    p_tr_head.text = "12-STAGE CONTROLLED PILOT EXECUTION TRACE"
    p_tr_head.font.name = FONT_HEADING
    p_tr_head.font.size = Pt(10)
    p_tr_head.font.bold = True
    p_tr_head.font.color.rgb = COLOR_ACCENT_CYAN

    trace_points = [
        "Reconstructed planar SECS boundary: 27.5 ft front × 40.0 ft depth.",
        "Deducted 3.0m road widening strip along southern road frontage.",
        "Applied Mumbai DCPR 2034 setbacks: 49.06 m² buildable envelope.",
        "CP-SAT solver generated 3 Pareto-optimal layouts; Option B selected.",
        "Compiled semantic CBM 3D model with 12/12 validation checks PASSED.",
        "Derived IS 1200 QTO: Concrete 29.45 m³, AAC Blockwork 38.12 m³.",
        "Computed 7-trade BOQ and regional cost waterfall using Q4 2026 rates.",
        "Persisted immutable release package with SHA-256 fingerprint."
    ]

    for pt in trace_points:
        p_pt = tf_tr.add_paragraph()
        p_pt.text = f"•  {pt}"
        p_pt.font.name = FONT_BODY
        p_pt.font.size = Pt(8.5)
        p_pt.font.color.rgb = COLOR_TEXT_BODY

    # Column 2: Verification Gate Status
    add_card(slide, Inches(6.813), y_col2, col2_w, h_col2, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)
    t_gate = slide.shapes.add_textbox(Inches(7.013), y_col2 + Inches(0.15), col2_w - Inches(0.4), h_col2 - Inches(0.3))
    tf_g = t_gate.text_frame
    tf_g.word_wrap = True
    tf_g.margin_left = tf_g.margin_top = tf_g.margin_right = tf_g.margin_bottom = 0

    p_g_head = tf_g.paragraphs[0]
    p_g_head.text = "PROFESSIONAL GATE AUDIT LEDGER (G0–G6)"
    p_g_head.font.name = FONT_HEADING
    p_g_head.font.size = Pt(10)
    p_g_head.font.bold = True
    p_g_head.font.color.rgb = COLOR_ACCENT_CYAN

    gates = [
        ("[G0] Site & Boundary Verification", "VERIFIED", "By Registered Surveyor SURV-MH-4421"),
        ("[G1] Regulatory Feasibility Review", "VERIFIED", "Mumbai DCPR 2034 setback & FSI compliance"),
        ("[G2] Structural Grid Review", "VERIFIED", "IS 456 RC column alignment & spans verified"),
        ("[G3] MEP & Services Review", "VERIFIED", "Vertical wet-wall stacking verified"),
        ("[G4] Cost & BOQ Review", "VERIFIED", "IS 1200 deductions & regional trade rates"),
        ("[G5] Municipal Build Authorization", "LOCKED", "Requires formal municipal building submission"),
        ("[G6] Contractor Construction Handoff", "LOCKED", "Requires executed commercial contract")
    ]

    for name, status, detail in gates:
        pg = tf_g.add_paragraph()
        pg.text = f"{name}:  {status} — {detail}"
        pg.font.name = FONT_BODY
        pg.font.size = Pt(8.5)
        if status == "VERIFIED":
            pg.font.color.rgb = COLOR_ACCENT_EMERALD
        else:
            pg.font.color.rgb = COLOR_ACCENT_AMBER

    # Footer note on pilot status
    n_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.35))
    tf_n = n_box.text_frame
    pn = tf_n.paragraphs[0]
    pn.text = "Controlled single-project pilot successfully executed. Municipal submission (G5) and construction handoff (G6) remain authority-controlled."
    pn.font.name = FONT_BODY
    pn.font.size = Pt(8.5)
    pn.font.color.rgb = COLOR_TEXT_SUBTLE


def build_slide_8(prs):
    """Slide 8: Trust / Engineering Foundation"""
    slide = create_base_slide(prs, "System Integrity", "Built for traceability, not just visualization.",
                             "AI assists with customer preferences; deterministic engineering algorithms and licensed professionals govern the truth.", 8)

    # Top Traceability Flow
    top_flow = add_card(slide, Inches(0.8), Inches(1.95), Inches(11.733), Inches(0.9), bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)
    t_tf = slide.shapes.add_textbox(Inches(1.0), Inches(2.05), Inches(11.333), Inches(0.7))
    tf = t_tf.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p0 = tf.paragraphs[0]
    p0.text = "THE DETERMINISTIC TRACEABILITY HIERARCHY"
    p0.font.name = FONT_HEADING
    p0.font.size = Pt(9)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT_CYAN

    p1 = tf.add_paragraph()
    p1.text = "Survey Evidence  →  Statutory Rules  →  Canonical Model (CBM)  →  Validation Checks  →  Quantities  →  Cost  →  Engineer Ledger  →  Release"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_TITLE

    # 2 Contrasting Comparison Panels
    col_w = Inches(5.72)
    gap = Inches(0.293)
    y_card = Inches(3.0)
    h_card = Inches(3.35)

    # Panel 1: Assisted Intelligence
    add_card(slide, Inches(0.8), y_card, col_w, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)
    t_ai = slide.shapes.add_textbox(Inches(1.05), y_card + Inches(0.2), col_w - Inches(0.5), h_card - Inches(0.4))
    tf_ai = t_ai.text_frame
    tf_ai.word_wrap = True
    tf_ai.margin_left = tf_ai.margin_top = tf_ai.margin_right = tf_ai.margin_bottom = 0

    p_ai_t = tf_ai.paragraphs[0]
    p_ai_t.text = "WHAT ASSISTIVE AI DOES"
    p_ai_t.font.name = FONT_HEADING
    p_ai_t.font.size = Pt(10)
    p_ai_t.font.bold = True
    p_ai_t.font.color.rgb = COLOR_ACCENT_CYAN

    p_ai_sub = tf_ai.add_paragraph()
    p_ai_sub.text = "Acceleration, natural input, and presentation"
    p_ai_sub.font.name = FONT_BODY
    p_ai_sub.font.size = Pt(11)
    p_ai_sub.font.bold = True
    p_ai_sub.font.color.rgb = COLOR_TEXT_TITLE

    tf_ai.add_paragraph()

    ai_bullets = [
        "Interprets natural language customer design briefs & preferences",
        "Assists with spatial adjacency suggestions and preliminary zoning",
        "Accelerates 3D visualization rendering and aesthetic styling",
        "Generates contextual summaries of complex statutory clauses"
    ]
    for b in ai_bullets:
        pb = tf_ai.add_paragraph()
        pb.text = f"•  {b}"
        pb.font.name = FONT_BODY
        pb.font.size = Pt(9.5)
        pb.font.color.rgb = COLOR_TEXT_MUTED

    # Panel 2: Invariants
    add_card(slide, Inches(6.813), y_card, col_w, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER_BLUE)
    t_inv = slide.shapes.add_textbox(Inches(7.063), y_card + Inches(0.2), col_w - Inches(0.5), h_card - Inches(0.4))
    tf_inv = t_inv.text_frame
    tf_inv.word_wrap = True
    tf_inv.margin_left = tf_inv.margin_top = tf_inv.margin_right = tf_inv.margin_bottom = 0

    p_inv_t = tf_inv.paragraphs[0]
    p_inv_t.text = "WHAT DETERMINISTIC RULES & ENGINEERS GOVERN"
    p_inv_t.font.name = FONT_HEADING
    p_inv_t.font.size = Pt(10)
    p_inv_t.font.bold = True
    p_inv_t.font.color.rgb = COLOR_ACCENT_EMERALD

    p_inv_sub = tf_inv.add_paragraph()
    p_inv_sub.text = "Absolute physical, legal & financial invariants"
    p_inv_sub.font.name = FONT_BODY
    p_inv_sub.font.size = Pt(11)
    p_inv_sub.font.bold = True
    p_inv_sub.font.color.rgb = COLOR_TEXT_TITLE

    tf_inv.add_paragraph()

    inv_bullets = [
        "Statutory setback distances, road reservations, and FSI caps",
        "Watertight 3D geometry topology and NBC 2016 room dimensions",
        "Structural column alignment, spans, and vertical wet-wall stacking",
        "Model-linked material takeoff calculations (IS 1200 deductions)",
        "Regional trade rate snapshots and construction cost waterfalls",
        "Licensed engineer sign-off and construction authorization (G0–G6)"
    ]
    for b in inv_bullets:
        pb = tf_inv.add_paragraph()
        pb.text = f"✓  {b}"
        pb.font.name = FONT_BODY
        pb.font.size = Pt(9.5)
        pb.font.color.rgb = COLOR_TEXT_BODY

    # Bottom Banner
    bot_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.45), Inches(11.733), Inches(0.4))
    tf_bot = bot_box.text_frame
    pbot = tf_bot.paragraphs[0]
    pbot.text = "Principle: AI never writes authoritative statutory setbacks, structural dimensions, quantities, or construction sign-offs."
    pbot.font.name = FONT_BODY
    pbot.font.size = Pt(9)
    pbot.font.bold = True
    pbot.font.color.rgb = COLOR_ACCENT_AMBER


def build_slide_9(prs):
    """Slide 9: Why Planwise Is Different"""
    slide = create_base_slide(prs, "Competitive Architecture", "Why Planwise is different",
                             "A structural comparison between Planwise, generative AI tools, and traditional fragmented consulting.", 9)

    # Comparison Table
    table_shape = slide.shapes.add_table(10, 4, Inches(0.8), Inches(2.05), Inches(11.733), Inches(4.35))
    table = table_shape.table

    # Column Widths
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(3.2)
    table.columns[2].width = Inches(2.6)
    table.columns[3].width = Inches(2.733)

    headers = ["Platform Dimension", "Planwise Enterprise", "Generative AI Tools", "Traditional Workflow"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(23, 37, 84) if j == 1 else RGBColor(19, 26, 38)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = FONT_HEADING
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_CYAN if j == 1 else COLOR_TEXT_TITLE

    rows = [
        ("Land Feasibility", "✓  Algorithmic statutory rulepack", "—  Ignored / No concept of land", "✓  Manual & slow (2–4 weeks)"),
        ("Rule-Aware Design", "✓  Multi-constraint CP-SAT solver", "—  Unconstrained prompt/pixels", "✓  Manual architectural drafting"),
        ("Canonical 3D Model", "✓  Single 3D source of truth (CBM)", "—  None (2D flat raster images)", "—  Fragmented CAD/BIM files"),
        ("Model-Linked QTO", "✓  Automated IS 1200 deductions", "—  None", "Manual spreadsheet takeoffs"),
        ("Traceable Costing", "✓  Live regional rate waterfall", "—  None", "Manual estimating / guesswork"),
        ("Professional Gates", "✓  Integrated G0–G6 audit ledger", "—  None", "Separate, delayed reviews"),
        ("Immutable Release", "✓  Cryptographic SHA-256 seal", "—  None", "—  Uncontrolled file versions"),
        ("Customer Experience", "✓  Transparent interactive web app", "✓  Simple text prompt", "—  Opaque technical documents"),
        ("Construction Path", "✓  Exportable IFC4 BIM & BOQ", "—  Pure artistic concept", "✓  Disjointed trade handoffs")
    ]

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.fill.solid()
            if j == 1:
                cell.fill.fore_color.rgb = RGBColor(17, 24, 39)
            else:
                cell.fill.fore_color.rgb = COLOR_CARD_BG if i % 2 == 0 else COLOR_CARD_ALT

            cell.text_frame.margin_left = Inches(0.1)
            cell.text_frame.margin_top = Inches(0.04)
            cell.text_frame.margin_bottom = Inches(0.04)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = FONT_BODY
            p.font.size = Pt(9)
            if j == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_TEXT_TITLE
            elif j == 1:
                p.font.bold = True
                p.font.color.rgb = COLOR_ACCENT_CYAN
            else:
                p.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom note
    n_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.35))
    tf_n = n_box.text_frame
    pn = tf_n.paragraphs[0]
    pn.text = "Planwise is not an image generator or CAD viewer; it is a full-stack computational development engine."
    pn.font.name = FONT_BODY
    pn.font.size = Pt(9)
    pn.font.bold = True
    pn.font.color.rgb = COLOR_ACCENT_AMBER


def build_slide_10(prs):
    """Slide 10: Business Model"""
    slide = create_base_slide(prs, "Commercial Strategy", "How Planwise can become a business",
                             "A pragmatic B2B / B2B2C business model combining transaction fees with recurring software subscriptions.", 10)

    streams = [
        ("01", "Per-Project Release Fee",
         "TRANSACTION / USAGE",
         "Tiered project release fee charged upon generating verified project packages, stamped IFC4 BIM models, and tender-ready BOQs.",
         ["Paid by landowners, builders, or project sponsors", "Low barrier to entry for one-off residential projects", "Scales directly with total development volume"]),
        ("02", "Professional SaaS",
         "B2B RECURRING SEAT",
         "Seat-based monthly/annual workflow subscription for architectural practices, structural engineers, and registered land surveyors.",
         ["Dedicated review ledgers & rulepack customizations", "Increases professional review throughput by 5–10x", "Sticky professional toolchain integration"]),
        ("03", "Builder Enterprise Tier",
         "B2B ANNUAL LICENSE",
         "Annual enterprise software license for residential builders and real-estate plotting developers managing ongoing project pipelines.",
         ["Unlimited parcel feasibility screenings", "Standardized house typology catalogues", "Direct procurement & supplier integration"]),
        ("04", "Computational API",
         "INFRASTRUCTURE LICENSING",
         "Headless API licensing for GIS portals, mortgage pre-qualification platforms, and property tech networks.",
         ["Instant setback, FSI, and cost estimation APIs", "High-volume, automated developer queries", "High-margin, low-touch software distribution"])
    ]

    w_card = Inches(2.78)
    gap = Inches(0.2)
    y_card = Inches(2.05)
    h_card = Inches(4.35)

    for i, (num, title, badge, desc, bullets) in enumerate(streams):
        x = Inches(0.8) + i * (w_card + gap)
        add_card(slide, x, y_card, w_card, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.2), y_card + Inches(0.2), w_card - Inches(0.4), h_card - Inches(0.4))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_num = tf.paragraphs[0]
        p_num.text = f"REVENUE STREAM {num}"
        p_num.font.name = FONT_HEADING
        p_num.font.size = Pt(8.5)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_ACCENT_CYAN

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_TITLE

        p_b = tf.add_paragraph()
        p_b.text = badge
        p_b.font.name = FONT_HEADING
        p_b.font.size = Pt(8)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_ACCENT_AMBER

        tf.add_paragraph()

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

        tf.add_paragraph()

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = f"•  {b}"
            pb.font.name = FONT_BODY
            pb.font.size = Pt(9)
            pb.font.color.rgb = COLOR_TEXT_BODY

    # Bottom note
    n_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.35))
    tf_n = n_box.text_frame
    pn = tf_n.paragraphs[0]
    pn.text = "Commercial model aligns incentives: Developers save pre-construction time; professionals expand billable review capacity."
    pn.font.name = FONT_BODY
    pn.font.size = Pt(8.5)
    pn.font.color.rgb = COLOR_TEXT_SUBTLE


def build_slide_11(prs):
    """Slide 11: Who Pays"""
    slide = create_base_slide(prs, "Customer Segments", "Who pays and why",
                             "A segmented breakdown of paying customers across the residential development ecosystem.", 11)

    segments = [
        ("PRIMARY CUSTOMERS", "Immediate Go-to-Market Focus", COLOR_ACCENT_CYAN,
         [
             ("Residential Builders", "Evaluate land parcels in hours instead of weeks; generate accurate initial cost envelopes before committing capital."),
             ("Property Developers", "Compress project kickoff cycles, eliminate manual re-drafting, and standardize multi-plot developments."),
             ("Architectural Practices", "Automate tedious statutory zoning and schematic drafting, focusing billable hours on high-value creative detailing."),
             ("Structural Engineering Firms", "Receive pre-coordinated geometric models with verified site boundaries, enabling efficient paid gate reviews.")
         ]),
        ("SECONDARY CUSTOMERS", "Growth Tier & Construction Partners", COLOR_ACCENT_BLUE,
         [
             ("General Contractors", "Receive verified 7-trade BOQs tied directly to 3D model elements, drastically reducing bid disputes and variation claims."),
             ("Real-Estate Advisory Firms", "Instant zoning, FSI, and feasibility assessments to accelerate high-value land brokerage transactions."),
             ("Housing Plotting Networks", "Scale standardized residential typologies across large plotted land subdivisions.")
         ]),
        ("FUTURE CUSTOMERS", "Enterprise Scale & Institutions", COLOR_ACCENT_EMERALD,
         [
             ("Proptech Portals & Networks", "Embed automated feasibility and development intelligence directly into property listing platforms."),
             ("Housing Authorities & Municipalities", "Automated regulatory compliance pre-checking for municipal building permit submissions."),
             ("Industrialized & Modular Builders", "Direct computational data feed into offsite pre-cast and modular manufacturing factories.")
         ])
    ]

    col_w = Inches(3.68)
    gap = Inches(0.34)
    y_card = Inches(2.05)
    h_card = Inches(4.35)

    for i, (title, subtitle, color, buyers) in enumerate(segments):
        x = Inches(0.8) + i * (col_w + gap)
        add_card(slide, x, y_card, col_w, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.2), y_card + Inches(0.2), col_w - Inches(0.4), h_card - Inches(0.4))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(9)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED

        tf.add_paragraph()

        for role, desc in buyers:
            pr = tf.add_paragraph()
            pr.text = f"{role}"
            pr.font.name = FONT_HEADING
            pr.font.size = Pt(10)
            pr.font.bold = True
            pr.font.color.rgb = COLOR_TEXT_TITLE

            pd = tf.add_paragraph()
            pd.text = desc
            pd.font.name = FONT_BODY
            pd.font.size = Pt(8.8)
            pd.font.color.rgb = COLOR_TEXT_MUTED

            tf.add_paragraph()


def build_slide_12(prs):
    """Slide 12: Business Value"""
    slide = create_base_slide(prs, "Value Proposition", "Why would a professional organization pay?",
                             "Measurable operational efficiency, risk mitigation, and transparency for every stakeholder.", 12)

    quads = [
        ("FOR THE BUILDER & DEVELOPER",
         "Capital Efficiency & Risk Reduction",
         COLOR_ACCENT_CYAN,
         [
             "Accelerates land evaluation from weeks to hours before committing land acquisition capital",
             "Eliminates early-stage cost blindspots with model-linked regional rate snapshots",
             "Standardizes pre-construction development workflows across distributed project teams",
             "Drastically reduces downstream change orders and construction contractor claims"
         ]),
        ("FOR THE ARCHITECTURAL PRACTICE",
         "Throughput & Workflow Modernization",
         COLOR_ACCENT_BLUE,
         [
             "Replaces manual setback drafting and statutory zoning calculations with automated rulepacks",
             "Generates mathematically sound starting baselines (CBM) ready for creative refinement",
             "Eliminates tedious re-drawing when clients modify basic spatial or budgetary preferences",
             "Expands firm design capacity without proportionally increasing drafting overhead"
         ]),
        ("FOR THE STRUCTURAL & MEP ENGINEER",
         "Traceability & Controlled Governance",
         COLOR_ACCENT_EMERALD,
         [
             "Receives pre-coordinated geometric models with verified geodetic site boundaries",
             "Enforces explicit professional review gates (G0–G6) backed by immutable audit trails",
             "Verifies RC column grids and vertical wet-wall stacking before detailed calculations",
             "Protects professional liability with clear, tamper-evident sign-off records"
         ]),
        ("FOR THE LANDOWNER & CUSTOMER",
         "Transparency & Informed Decisions",
         COLOR_ACCENT_AMBER,
         [
             "Demystifies complex municipal bylaws and construction planning into clear choices",
             "Guaranteed consistency: The selected floor plan strictly drives material quantities and cost",
             "Avoids contractor budget shocks by establishing an audited preliminary cost baseline",
             "Provides a clear, verified roadmap to licensed construction execution"
         ])
    ]

    w_card = Inches(5.72)
    gap = Inches(0.293)
    h_card = Inches(2.1)

    for i, (title, subtitle, color, bullets) in enumerate(quads):
        col = i % 2
        row = i // 2
        x = Inches(0.8) + col * (w_card + gap)
        y = Inches(2.05) + row * (h_card + Inches(0.25))

        add_card(slide, x, y, w_card, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.25), y + Inches(0.15), w_card - Inches(0.5), h_card - Inches(0.3))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(9)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_TEXT_MUTED

        tf.add_paragraph()

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = f"•  {b}"
            pb.font.name = FONT_BODY
            pb.font.size = Pt(8.8)
            pb.font.color.rgb = COLOR_TEXT_BODY


def build_slide_13(prs):
    """Slide 13: Technology Moat"""
    slide = create_base_slide(prs, "Competitive Moat", "The moat is the system, not the AI image.",
                             "Defensibility is rooted in the tightly coupled, version-controlled computational graph linking physical land to cost.", 13)

    # Top System Interlocking Box
    top_box = add_card(slide, Inches(0.8), Inches(1.95), Inches(11.733), Inches(1.05), bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER_BLUE)
    t_top = slide.shapes.add_textbox(Inches(1.0), Inches(2.05), Inches(11.333), Inches(0.85))
    tf_top = t_top.text_frame
    tf_top.margin_left = tf_top.margin_top = tf_top.margin_right = tf_top.margin_bottom = 0

    pt0 = tf_top.paragraphs[0]
    pt0.text = "THE CONNECTED COMPUTATIONAL GRAPH"
    pt0.font.name = FONT_HEADING
    pt0.font.size = Pt(9)
    pt0.font.bold = True
    pt0.font.color.rgb = COLOR_ACCENT_CYAN

    pt1 = tf_top.add_paragraph()
    pt1.text = "Site Intelligence  +  Statutory Rules  +  CP-SAT Solvers  +  Canonical 3D Model (CBM)  +  IS 1200 QTO  +  Cost Engine  +  Review Ledger"
    pt1.font.name = FONT_HEADING
    pt1.font.size = Pt(11)
    pt1.font.bold = True
    pt1.font.color.rgb = COLOR_TEXT_TITLE

    pt2 = tf_top.add_paragraph()
    pt2.text = "A change in any layer propagates mathematically across every downstream artifact in milliseconds."
    pt2.font.name = FONT_BODY
    pt2.font.size = Pt(9)
    pt2.font.color.rgb = COLOR_TEXT_MUTED

    # 3 Moat Pillars
    col_w = Inches(3.68)
    gap = Inches(0.34)
    y_card = Inches(3.2)
    h_card = Inches(3.2)

    pillars = [
        ("01", "Coordinated Change Propagation",
         "Physical & mathematical binding",
         [
             "A 0.5m modification to a site boundary instantly cascades through statutory setbacks, CP-SAT room envelopes, 3D geometry, concrete quantities, and cost waterfalls.",
             "Standalone CAD, BIM, or estimating software cannot propagate changes across disciplines automatically.",
             "Prevents costly human errors and synchronization gaps."
         ]),
        ("02", "Deterministic Integrity",
         "Auditable engineering proofs",
         [
             "Unlike probabilistic AI image generators that hallucinate impossible layouts, Planwise relies on formal constraint satisfaction (Google OR-Tools).",
             "Every calculation produces an auditable geometric and arithmetic proof suitable for licensed engineers.",
             "Creates legal and regulatory defensibility that builds institutional trust."
         ]),
        ("03", "Institutional Rulepack Moat",
         "Accumulating regional intelligence",
         [
             "Local municipal bylaws (e.g. Mumbai DCPR 2034) and regional trade rate snapshots build an expanding institutional data barrier.",
             "Expanding rulepacks create high switching costs once builders and developers standardize on Planwise.",
             "Network effects: More verified pilot projects enhance regional rate and rule accuracy."
         ])
    ]

    for i, (num, title, sub, bullets) in enumerate(pillars):
        x = Inches(0.8) + i * (col_w + gap)
        add_card(slide, x, y_card, col_w, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.2), y_card + Inches(0.18), col_w - Inches(0.4), h_card - Inches(0.36))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_num = tf.paragraphs[0]
        p_num.text = f"DEFENSIBILITY PILLAR {num}"
        p_num.font.name = FONT_HEADING
        p_num.font.size = Pt(8.5)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_ACCENT_CYAN

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_TITLE

        p_sub = tf.add_paragraph()
        p_sub.text = sub
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(9)
        p_sub.font.color.rgb = COLOR_ACCENT_AMBER

        tf.add_paragraph()

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = f"•  {b}"
            pb.font.name = FONT_BODY
            pb.font.size = Pt(8.8)
            pb.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom note
    n_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.35))
    tf_n = n_box.text_frame
    pn = tf_n.paragraphs[0]
    pn.text = "Defensibility comes from connecting these layers into one versioned, auditable workflow."
    pn.font.name = FONT_BODY
    pn.font.size = Pt(8.5)
    pn.font.color.rgb = COLOR_TEXT_SUBTLE


def build_slide_14(prs):
    """Slide 14: Roadmap"""
    slide = create_base_slide(prs, "Platform Evolution", "Platform roadmap: From pilot to enterprise scale",
                             "Disciplined technological execution transitioning from core engine verification to multi-market expansion.", 14)

    phases = [
        ("PHASE 1: NOW", "Implemented & Verified", COLOR_ACCENT_CYAN,
         [
             "SECS planar land intake & geodetic GPS anchoring",
             "Mumbai DCPR 2034 statutory feasibility rule engine",
             "CP-SAT parametric 2D/3D house layout generator",
             "Canonical Building Model (CBM) semantic 3D schema",
             "Model-linked QTO takeoff (IS 1200 measurement rules)",
             "Regional cost waterfall engine (INDIA-MUMBAI-2026-Q4-V1)",
             "G0–G6 Engineer Review Ledger & SHA-256 release package",
             "Controlled 1,100 sq ft single-project pilot completed"
         ]),
        ("PHASE 2: NEXT", "Professional Pilots & Ecosystem", COLOR_ACCENT_BLUE,
         [
             "Interactive client walkthroughs & 3D WebGL model inspection",
             "Multi-firm professional pilot cohort (builders, architects, engineers)",
             "Municipal pre-check format exports (AutoDCR / CAD compliance)",
             "Subcontractor procurement packages & digital tender generation",
             "Mobile surveyor intake app with timestamped photo anchoring",
             "Expanded regional cost database for Tier-1 Indian metros"
         ]),
        ("PHASE 3: LATER", "Multi-Typology & Enterprise Scale", COLOR_ACCENT_EMERALD,
         [
             "Multi-typology expansion: commercial, mixed-use, row housing",
             "National statutory rulepack library (NBC 2016, Delhi, Bengaluru)",
             "Automated structural FEA integration & reinforcement schedules",
             "On-site construction IoT progress tracking & digital twin syncing",
             "Enterprise ERP integration for large real-estate development corporations"
         ])
    ]

    col_w = Inches(3.68)
    gap = Inches(0.34)
    y_card = Inches(2.05)
    h_card = Inches(4.35)

    for i, (title, subtitle, color, items) in enumerate(phases):
        x = Inches(0.8) + i * (col_w + gap)
        add_card(slide, x, y_card, col_w, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.2), y_card + Inches(0.2), col_w - Inches(0.4), h_card - Inches(0.4))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(9.5)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_TEXT_MUTED

        tf.add_paragraph()

        for item in items:
            pi = tf.add_paragraph()
            bullet_sym = "✓" if "NOW" in title else "→"
            pi.text = f"{bullet_sym}  {item}"
            pi.font.name = FONT_BODY
            pi.font.size = Pt(9)
            if "NOW" in title:
                pi.font.color.rgb = COLOR_TEXT_BODY
            else:
                pi.font.color.rgb = COLOR_TEXT_MUTED


def build_slide_15(prs):
    """Slide 15: What We Are Looking For"""
    slide = create_base_slide(prs, "Partnership Opportunities", "We are looking for the right partners.",
                             "Collaborating with forward-thinking builders, professionals, and strategic investors to scale computational development.", 15)

    asks = [
        ("01", "Builder & Developer Pilots",
         "PILOT COLLABORATIONS",
         COLOR_ACCENT_CYAN,
         "Forward-thinking residential builders and plotting developers seeking to pilot Planwise on upcoming real estate parcels for automated feasibility and costing."),
        ("02", "Professional Review Network",
         "ARCHITECTS & ENGINEERS",
         COLOR_ACCENT_BLUE,
         "Registered architectural practices, certified structural engineers, and licensed land surveyors interested in joining our accredited review network."),
        ("03", "Real Pilot Parcels",
         "PROJECT SITES",
         COLOR_ACCENT_EMERALD,
         "Landowners, builders, and plotting developers with residential plots (800–2,500 sq ft) ready for feasibility screening, layout generation, and costing."),
        ("04", "Technology & Industry Partners",
         "ECOSYSTEM INTEGRATION",
         COLOR_ACCENT_AMBER,
         "Collaborations with GIS data providers, BIM software ecosystems, municipal compliance systems, and construction technology platforms."),
        ("05", "Strategic Capital & Advisors",
         "INVESTMENT & GUIDANCE",
         RGBColor(168, 85, 247), # Purple
         "Value-add strategic investors and real-estate industry leaders aligned with modernizing fragmented property development into a scalable computational science.")
    ]

    w_card = Inches(2.2)
    gap = Inches(0.183)
    y_card = Inches(2.05)
    h_card = Inches(4.35)

    for i, (num, title, badge, color, desc) in enumerate(asks):
        x = Inches(0.8) + i * (w_card + gap)
        add_card(slide, x, y_card, w_card, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.15), y_card + Inches(0.18), w_card - Inches(0.3), h_card - Inches(0.36))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_num = tf.paragraphs[0]
        p_num.text = f"OPPORTUNITY {num}"
        p_num.font.name = FONT_HEADING
        p_num.font.size = Pt(8.5)
        p_num.font.bold = True
        p_num.font.color.rgb = color

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_TITLE

        p_b = tf.add_paragraph()
        p_b.text = badge
        p_b.font.name = FONT_BODY
        p_b.font.size = Pt(7.5)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_TEXT_MUTED

        tf.add_paragraph()

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_BODY

    # Bottom note
    n_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.35))
    tf_n = n_box.text_frame
    pn = tf_n.paragraphs[0]
    pn.text = "Focus: Quality pilot partnerships with active parcels and accredited professionals."
    pn.font.name = FONT_BODY
    pn.font.size = Pt(8.5)
    pn.font.color.rgb = COLOR_TEXT_SUBTLE


def build_slide_16(prs):
    """Slide 16: The Ask / Next Step"""
    slide = create_base_slide(prs, "Call to Action", "Let's build the next generation of development infrastructure.",
                             "Transforming fragmented real estate development into an integrated, computational science.", 16)

    # 3 Action Pathways
    col_w = Inches(3.68)
    gap = Inches(0.34)
    y_card = Inches(2.05)
    h_card = Inches(2.9)

    actions = [
        ("01", "Pilot Planwise",
         "FOR BUILDERS & LANDOWNERS",
         COLOR_ACCENT_CYAN,
         "Bring an active residential parcel to experience the end-to-end feasibility, design, QTO, and costing workflow on real site data.",
         "Schedule a 30-Minute Pilot Demonstration"),
        ("02", "Join the Network",
         "FOR ARCHITECTS & ENGINEERS",
         COLOR_ACCENT_BLUE,
         "Become a founding professional partner in our accredited review network for statutory compliance and structural sign-offs.",
         "Apply for Accredited Partner Network"),
        ("03", "Invest & Scale",
         "FOR STRATEGIC INVESTORS",
         COLOR_ACCENT_EMERALD,
         "Partner with Planwise Enterprise as a strategic investor to modernise the residential development and construction software ecosystem.",
         "Request Detailed Investment Brief")
    ]

    for i, (num, title, badge, color, desc, cta) in enumerate(actions):
        x = Inches(0.8) + i * (col_w + gap)
        add_card(slide, x, y_card, col_w, h_card, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        t_box = slide.shapes.add_textbox(x + Inches(0.2), y_card + Inches(0.18), col_w - Inches(0.4), h_card - Inches(0.36))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_num = tf.paragraphs[0]
        p_num.text = f"PATHWAY {num}"
        p_num.font.name = FONT_HEADING
        p_num.font.size = Pt(8.5)
        p_num.font.bold = True
        p_num.font.color.rgb = color

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_TITLE

        p_b = tf.add_paragraph()
        p_b.text = badge
        p_b.font.name = FONT_BODY
        p_b.font.size = Pt(8)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_TEXT_MUTED

        tf.add_paragraph()

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_BODY

        tf.add_paragraph()

        p_cta = tf.add_paragraph()
        p_cta.text = f"→  {cta}"
        p_cta.font.name = FONT_HEADING
        p_cta.font.size = Pt(9.5)
        p_cta.font.bold = True
        p_cta.font.color.rgb = color

    # Bottom Contact & Organization Box
    bot_card = add_card(slide, Inches(0.8), Inches(5.15), Inches(11.733), Inches(1.5), bg_color=RGBColor(23, 37, 84), border_color=COLOR_ACCENT_BLUE)

    t_bot = slide.shapes.add_textbox(Inches(1.1), Inches(5.25), Inches(11.133), Inches(1.3))
    tf_b = t_bot.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0

    pb0 = tf_b.paragraphs[0]
    pb0.text = "PLANWISE ENTERPRISE"
    pb0.font.name = FONT_HEADING
    pb0.font.size = Pt(14)
    pb0.font.bold = True
    pb0.font.color.rgb = COLOR_TEXT_TITLE

    pb1 = tf_b.add_paragraph()
    pb1.text = "Computational Real Estate Development & Construction Platform"
    pb1.font.name = FONT_BODY
    pb1.font.size = Pt(10.5)
    pb1.font.color.rgb = COLOR_ACCENT_CYAN

    pb2 = tf_b.add_paragraph()
    pb2.text = "Location: Mumbai, Maharashtra, India  •  Audience Outreach: Developers, Builders, AEC Professionals & Strategic Investors"
    pb2.font.name = FONT_BODY
    pb2.font.size = Pt(9.5)
    pb2.font.color.rgb = COLOR_TEXT_MUTED

    pb3 = tf_b.add_paragraph()
    pb3.text = "Website: planwise.enterprise  •  Inquiries: partnerships@planwise.enterprise  •  Private & Confidential"
    pb3.font.name = FONT_BODY
    pb3.font.size = Pt(9.5)
    pb3.font.bold = True
    pb3.font.color.rgb = COLOR_TEXT_BODY


# =========================================================
# MAIN GENERATOR
# =========================================================

def main():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    print("Building Slide 1: Cover...")
    build_slide_1(prs)

    print("Building Slide 2: The Problem...")
    build_slide_2(prs)

    print("Building Slide 3: The Planwise Solution...")
    build_slide_3(prs)

    print("Building Slide 4: Product Modules...")
    build_slide_4(prs)

    print("Building Slide 5: Customer Experience...")
    build_slide_5(prs)

    print("Building Slide 6: What Is Already Built...")
    build_slide_6(prs)

    print("Building Slide 7: Real Project / Pilot...")
    build_slide_7(prs)

    print("Building Slide 8: Trust / Engineering Foundation...")
    build_slide_8(prs)

    print("Building Slide 9: Why Planwise Is Different...")
    build_slide_9(prs)

    print("Building Slide 10: Business Model...")
    build_slide_10(prs)

    print("Building Slide 11: Who Pays...")
    build_slide_11(prs)

    print("Building Slide 12: Business Value...")
    build_slide_12(prs)

    print("Building Slide 13: Technology Moat...")
    build_slide_13(prs)

    print("Building Slide 14: Roadmap...")
    build_slide_14(prs)

    print("Building Slide 15: What We Are Looking For...")
    build_slide_15(prs)

    print("Building Slide 16: The Ask / Next Step...")
    build_slide_16(prs)

    output_pptx = "PLANWISE_BUSINESS_PITCH_DECK.pptx"
    prs.save(output_pptx)
    print(f"Successfully generated {output_pptx} with 16 slides!")


if __name__ == "__main__":
    main()
